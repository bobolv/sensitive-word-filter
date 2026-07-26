from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from .engine import SensitiveWordFilter


def _is_scannable(value) -> bool:
    return isinstance(value, str) and not value.startswith("=")


def scan_xlsx(
    content: bytes,
    engine: SensitiveWordFilter,
    min_level: int = 1,
    preview_limit: int = 200_000,
    match_limit: int = 2_000,
) -> dict:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=True)
    matches: list[dict] = []
    total_count = 0
    total_chars = 0
    preview_parts: list[str] = []
    preview_chars = 0

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not _is_scannable(value):
                        continue
                    total_chars += len(value)
                    if preview_chars < preview_limit:
                        line = f"[{worksheet.title}!{cell.coordinate}] {value}\n"
                        remaining = preview_limit - preview_chars
                        preview_parts.append(line[:remaining])
                        preview_chars += min(len(line), remaining)
                    result = engine.scan(value, min_level)
                    total_count += len(result.matches)
                    for match in result.matches:
                        if len(matches) >= match_limit:
                            break
                        item = match.to_dict()
                        item.update({"sheet": worksheet.title, "cell": cell.coordinate})
                        matches.append(item)
    finally:
        workbook.close()

    preview = "".join(preview_parts)
    return {
        "sensitive": total_count > 0,
        "count": total_count,
        "matches": matches,
        "matches_truncated": total_count > len(matches),
        "text": preview,
        "text_preview": preview,
        "text_length": total_chars,
        "text_truncated": total_chars > preview_limit,
    }


def replace_xlsx(
    content: bytes,
    engine: SensitiveWordFilter,
    replacement: str = "*",
    min_level: int = 1,
):
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=True)
    changed_cells = 0
    total_matches = 0

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if not _is_scannable(cell.value):
                    continue
                output, result = engine.replace(cell.value, replacement, min_level)
                if output != cell.value:
                    cell.value = output
                    changed_cells += 1
                    total_matches += len(result.matches)

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), {
        "sensitive": total_matches > 0,
        "count": total_matches,
        "changed_cells": changed_cells,
    }
