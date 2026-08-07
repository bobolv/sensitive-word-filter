from __future__ import annotations

from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .engine import SensitiveWordFilter


HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*([1-9])$", re.IGNORECASE)
HEADING_STYLE_ID = re.compile(r"^Heading([1-9])$", re.IGNORECASE)
FIGURE_CAPTION = re.compile(
    r"^\s*(?:图|Figure)\s*(?:\d+(?:[.-]\d+)*|[一二三四五六七八九十百]+)",
    re.IGNORECASE,
)
MANUAL_HEADING = re.compile(
    r"^\s*(?:"
    r"第[一二三四五六七八九十百千万0-9]+[章节篇编部分]\s*|"
    r"[一二三四五六七八九十百]+[、．.]\s*|"
    r"[（(][一二三四五六七八九十百0-9]+[）)]\s*|"
    r"\d+(?:\.\d+){1,5}[、．.\s]*|"
    r"\d+[、．.]\s*"
    r")"
)


def _part_paragraphs(element, owner, seen: set[object]):
    """遍历部件内所有段落，包括表格、嵌套表格和内容控件。"""
    for paragraph_element in element.iter(qn("w:p")):
        if paragraph_element in seen:
            continue
        seen.add(paragraph_element)
        yield Paragraph(paragraph_element, owner)


def iter_document_paragraphs(document):
    seen: set[object] = set()
    yield from _part_paragraphs(document.element.body, document, seen)

    seen_parts: set[str] = set()
    for section in document.sections:
        for part in (section.header, section.footer):
            key = str(part.part.partname)
            if key in seen_parts:
                continue
            seen_parts.add(key)
            yield from _part_paragraphs(part._element, part, seen)


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    """包含普通、超链接及域结果中的所有文字运行。"""
    return [Run(element, paragraph) for element in paragraph._p.iter(qn("w:r"))]


def _paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in _paragraph_runs(paragraph))


def document_text(document) -> str:
    texts = (_paragraph_text(paragraph) for paragraph in iter_document_paragraphs(document))
    return "\n".join(text for text in texts if text)


def _body_items(document):
    """按 Word 主体中的实际顺序遍历段落和表格。"""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None:
        outline = paragraph_properties.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                return min(int(outline.get(qn("w:val"))) + 1, 6)
            except (TypeError, ValueError):
                pass

    # 很多业务文档使用“基于标题样式”的自定义样式，标题级别保存在样式
    # 或其基准样式中，而不是直接写在段落属性中。
    style = paragraph.style
    seen_styles: set[str] = set()
    while style is not None and style.style_id not in seen_styles:
        seen_styles.add(style.style_id)
        style_name = style.name.strip()
        match = HEADING_STYLE.match(style_name) or HEADING_STYLE_ID.match(style.style_id)
        if match:
            return min(int(match.group(1)), 6)
        if style_name.casefold() in {"title", "标题"}:
            return 1
        style_properties = style.element.pPr
        if style_properties is not None:
            outline = style_properties.find(qn("w:outlineLvl"))
            if outline is not None:
                try:
                    return min(int(outline.get(qn("w:val"))) + 1, 6)
                except (TypeError, ValueError):
                    pass
        style = style.base_style
    return None


def _is_toc_paragraph(paragraph: Paragraph) -> bool:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    normalized_style = style_name.strip().casefold()
    if normalized_style.startswith("toc") or normalized_style.startswith("目录"):
        return True
    instructions = " ".join(
        node.text or "" for node in paragraph._p.iter(qn("w:instrText"))
    ).upper()
    return "TOC " in instructions


def _normalized_title_text(value: str) -> str:
    value = value.replace("\t", " ").strip()
    value = re.sub(r"[.．·…\s]+\d+\s*$", "", value)
    return re.sub(r"\s+", "", value)


def _toc_title_texts(document) -> set[str]:
    titles: set[str] = set()
    for paragraph in iter_document_paragraphs(document):
        if _is_toc_paragraph(paragraph):
            normalized = _normalized_title_text(_paragraph_text(paragraph))
            if normalized:
                titles.add(normalized)
    return titles


def _effective_style_value(paragraph: Paragraph, attribute: str):
    direct_value = getattr(paragraph.paragraph_format, attribute, None)
    if direct_value is not None:
        return direct_value
    style = paragraph.style
    seen_styles: set[str] = set()
    while style is not None and style.style_id not in seen_styles:
        seen_styles.add(style.style_id)
        value = getattr(style.paragraph_format, attribute, None)
        if value is not None:
            return value
        style = style.base_style
    return None


def _largest_font_size(paragraph: Paragraph) -> float:
    sizes = [run.font.size.pt for run in paragraph.runs if run.font.size is not None]
    style = paragraph.style
    seen_styles: set[str] = set()
    while style is not None and style.style_id not in seen_styles:
        seen_styles.add(style.style_id)
        if style.font.size is not None:
            sizes.append(style.font.size.pt)
        style = style.base_style
    return max(sizes, default=0)


def _looks_like_visual_heading(paragraph: Paragraph, toc_titles: set[str]) -> bool:
    text = paragraph.text.strip()
    if not text or len(text) > 200:
        return False
    if _normalized_title_text(text) in toc_titles:
        return True
    if MANUAL_HEADING.match(text) and not text.endswith(("。", "；", ";")):
        return True
    alignment = _effective_style_value(paragraph, "alignment")
    nonempty_runs = [run for run in paragraph.runs if run.text.strip()]
    all_bold = bool(nonempty_runs) and all(run.bold is True for run in nonempty_runs)
    return (
        alignment == WD_ALIGN_PARAGRAPH.CENTER
        and len(text) <= 100
        and not text.endswith(("。", "；", ";", "，", ","))
        and (_largest_font_size(paragraph) >= 14 or all_bold)
    )


def _markdown_text(value: str) -> str:
    return (
        value.replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\n", "<br>")
        .strip()
    )


def _table_markdown(table: Table) -> list[str]:
    rows: list[list[str]] = []
    seen_cells: set[object] = set()
    for row in table.rows:
        values: list[str] = []
        for cell in row.cells:
            cell_key = cell._tc
            if cell_key in seen_cells:
                values.append("")
            else:
                seen_cells.add(cell_key)
                values.append(_markdown_text(cell.text))
        rows.append(values)
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return lines


def docx_to_markdown(
    content: bytes,
    engine: SensitiveWordFilter,
    replacement: str = "*",
    min_level: int = 1,
) -> tuple[bytes, dict]:
    """先执行敏感词替换，再将标题、正文和表格转换为 Markdown。"""
    filtered_content, replace_result = replace_docx(content, engine, replacement, min_level)
    document = Document(BytesIO(filtered_content))
    blocks: list[str] = []
    heading_count = 0
    table_count = 0

    for item in _body_items(document):
        if isinstance(item, Paragraph):
            if _is_toc_paragraph(item):
                continue
            value = item.text.strip()
            if not value:
                continue
            level = _heading_level(item)
            if level is not None:
                blocks.append(f"{'#' * level} {value}")
                heading_count += 1
            else:
                blocks.append(value)
        else:
            markdown_table = _table_markdown(item)
            if markdown_table:
                blocks.append("\n".join(markdown_table))
                table_count += 1

    markdown = "\n\n".join(blocks).rstrip() + "\n"
    return b"\xef\xbb\xbf" + markdown.encode("utf-8"), {
        **replace_result,
        "heading_count": heading_count,
        "table_count": table_count,
    }


def _remove_drawings(paragraph: Paragraph) -> None:
    for drawing in list(paragraph._p.iter(qn("w:drawing"))):
        parent = drawing.getparent()
        if parent is not None:
            parent.remove(drawing)
    for picture in list(paragraph._p.iter(qn("w:pict"))):
        parent = picture.getparent()
        if parent is not None:
            parent.remove(picture)


def _is_figure_caption(paragraph: Paragraph) -> bool:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    return style_name.strip().casefold() in {"caption", "题注"} or bool(
        FIGURE_CAPTION.match(paragraph.text)
    )


def _clear_cell_except_nested_tables(cell) -> None:
    cell_properties = cell._tc.tcPr
    for child in list(cell._tc):
        if child is cell_properties:
            continue
        if child.tag == qn("w:tbl"):
            _clear_table_body(Table(child, cell))
        else:
            cell._tc.remove(child)
    cell._tc.append(OxmlElement("w:p"))


def _clear_table_body(table: Table) -> None:
    """保留首行表头（及显式重复表头行），清空其余单元格内容。"""
    if not table.rows:
        return
    header_rows = 1
    for row in table.rows[1:]:
        row_properties = row._tr.trPr
        repeated_header = (
            row_properties is not None
            and row_properties.find(qn("w:tblHeader")) is not None
        )
        if not repeated_header:
            break
        header_rows += 1

    seen_cells: set[object] = set()
    for row in table.rows[header_rows:]:
        for cell in row.cells:
            cell_key = cell._tc
            if cell_key in seen_cells:
                continue
            seen_cells.add(cell_key)
            _clear_cell_except_nested_tables(cell)


def create_writing_template(content: bytes) -> tuple[bytes, dict]:
    """保留标题、表头、表格结构和图片题注，生成精简 Word 模板。"""
    document = Document(BytesIO(content))
    body = document.element.body
    toc_titles = _toc_title_texts(document)
    heading_count = 0
    table_count = 0
    caption_count = 0

    for child in list(body.iterchildren()):
        if child.tag == qn("w:sectPr"):
            continue
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            if _is_toc_paragraph(paragraph):
                body.remove(child)
                continue
            if (
                paragraph.text.strip()
                and (
                    _heading_level(paragraph) is not None
                    or _looks_like_visual_heading(paragraph, toc_titles)
                )
            ):
                _remove_drawings(paragraph)
                heading_count += 1
            elif _is_figure_caption(paragraph) and paragraph.text.strip():
                _remove_drawings(paragraph)
                caption_count += 1
            else:
                body.remove(child)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            _clear_table_body(table)
            # Word 会把直接相邻且格式兼容的表格自动连成一张；保留一个空段落隔离。
            child.addnext(OxmlElement("w:p"))
            table_count += 1
        else:
            body.remove(child)

    output = BytesIO()
    document.save(output)
    return output.getvalue(), {
        "heading_count": heading_count,
        "table_count": table_count,
        "caption_count": caption_count,
    }


def _limited_result(result, match_limit: int) -> dict:
    return {
        "sensitive": result.sensitive,
        "count": len(result.matches),
        "matches": [match.to_dict() for match in result.matches[:match_limit]],
        "matches_truncated": len(result.matches) > match_limit,
    }


def scan_docx(
    content: bytes,
    engine: SensitiveWordFilter,
    min_level: int = 1,
    preview_limit: int = 200_000,
    match_limit: int = 2_000,
) -> dict:
    text = document_text(Document(BytesIO(content)))
    result = engine.scan(text, min_level)
    preview = text[:preview_limit]
    return {
        "text": preview,
        "text_preview": preview,
        "text_length": len(text),
        "text_truncated": len(text) > preview_limit,
        **_limited_result(result, match_limit),
    }


def _replace_range(paragraph, start: int, end: int, replacement: str) -> None:
    positions, cursor = [], 0
    runs = _paragraph_runs(paragraph)
    for run in runs:
        positions.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    touched = []
    for index, (run_start, run_end) in enumerate(positions):
        local_start = max(start, run_start) - run_start
        local_end = min(end, run_end) - run_start
        if local_start < local_end:
            touched.append((index, local_start, local_end))
    if not touched:
        return

    first = touched[0][0]
    for index, local_start, local_end in reversed(touched):
        run = runs[index]
        inserted = replacement if index == first else ""
        run.text = run.text[:local_start] + inserted + run.text[local_end:]


def replace_docx(
    content: bytes,
    engine: SensitiveWordFilter,
    replacement: str = "*",
    min_level: int = 1,
):
    document = Document(BytesIO(content))
    total_matches = 0
    changed_paragraphs = 0

    # 按段处理，避免大文档再次拼接一份完整文本用于替换。
    for paragraph in iter_document_paragraphs(document):
        ranges = engine.replacement_ranges(_paragraph_text(paragraph), replacement, min_level)
        total_matches += len(ranges)
        if ranges:
            changed_paragraphs += 1
        for start, end, value in reversed(ranges):
            _replace_range(paragraph, start, end, value)

    output = BytesIO()
    document.save(output)
    return output.getvalue(), {
        "sensitive": total_matches > 0,
        "count": total_matches,
        "changed_paragraphs": changed_paragraphs,
    }
