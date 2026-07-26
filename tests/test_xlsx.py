from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from sensitive_filter import SensitiveWordFilter
from sensitive_filter.xlsx_handler import replace_xlsx, scan_xlsx


engine = SensitiveWordFilter([
    {"word": "张三", "category": "姓名", "level": 1, "replacement": "某人员"},
    {"word": "机密", "category": "保密", "level": 2},
])

workbook = Workbook()
sheet = workbook.active
sheet.title = "数据表"
sheet["A1"] = "张三查看机密"
sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
sheet["B1"] = "=1+1"
sheet.merge_cells("C1:D1")
sheet["C1"] = "机密内容"
hidden = workbook.create_sheet("隐藏页")
hidden.sheet_state = "hidden"
hidden["A1"] = "张三"
source = BytesIO()
workbook.save(source)

scanned = scan_xlsx(source.getvalue(), engine)
assert scanned["count"] == 4
assert any(item["sheet"] == "数据表" and item["cell"] == "A1" for item in scanned["matches"])
assert any(item["sheet"] == "隐藏页" for item in scanned["matches"])

output, result = replace_xlsx(source.getvalue(), engine, "*")
assert result["count"] == 4
assert result["changed_cells"] == 3
filtered = load_workbook(BytesIO(output), data_only=False)
assert filtered["数据表"]["A1"].value == "某人员查看**"
assert filtered["数据表"]["B1"].value == "=1+1"
assert filtered["数据表"]["C1"].value == "**内容"
assert filtered["隐藏页"]["A1"].value == "某人员"
assert filtered["数据表"]["A1"].fill.fgColor.rgb == "00FFFF00"
assert str(filtered["数据表"].merged_cells) == "C1:D1"

reverse_engine = SensitiveWordFilter([
    {"word": "某人员", "replacement": "张三", "category": "姓名", "level": 1},
    {"word": "**", "replacement": "机密", "category": "保密", "level": 1},
])
restored_output, restored_result = replace_xlsx(output, reverse_engine)
restored = load_workbook(BytesIO(restored_output), data_only=False)
assert restored_result["count"] == 4
assert restored["数据表"]["A1"].value == "张三查看机密"
assert restored["数据表"]["C1"].value == "机密内容"
print("XLSX multi-sheet/formula/style/merged scan and replace passed")
