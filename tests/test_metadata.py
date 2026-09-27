import os
import socket
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook, load_workbook
from lxml import etree
from openpyxl.comments import Comment
from sensitive_filter import SensitiveWordFilter

from sensitive_filter.batch import replace_batch
from sensitive_filter.metadata import clear_author_properties


@unittest.skipUnless(os.name == "nt", "Windows file timestamps")
class MetadataTests(unittest.TestCase):
    def test_review_history_and_embedded_payload(self):
        document = Document()
        document.add_paragraph("保留正文")
        stream = BytesIO()
        document.save(stream)
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        altered = BytesIO()
        with ZipFile(BytesIO(stream.getvalue())) as original, ZipFile(altered, "w") as target:
            for info in original.infolist():
                data = original.read(info)
                if info.filename == "word/document.xml":
                    root = etree.fromstring(data)
                    paragraph = root.find(".//" + namespace + "p")
                    for tag, text in [("ins", "新增正文"), ("del", "已删除的个人秘密")]:
                        revision = etree.SubElement(paragraph, namespace + tag)
                        revision.set(namespace + "author", "修订者姓名")
                        run = etree.SubElement(revision, namespace + "r")
                        etree.SubElement(run, namespace + ("t" if tag == "ins" else "delText")).text = text
                    data = etree.tostring(root)
                target.writestr(info, data)
        output, report = clear_author_properties(altered.getvalue(), ".docx")
        self.assertGreater(report["count"], 0)
        self.assertEqual(Document(BytesIO(output)).paragraphs[0].text, "保留正文新增正文")
        with ZipFile(BytesIO(output)) as archive:
            xml = archive.read("word/document.xml").decode()
            self.assertNotIn("修订者姓名", xml)
            self.assertNotIn("已删除的个人秘密", xml)
        embedded = BytesIO(output)
        with ZipFile(embedded, "a") as archive:
            archive.writestr("word/embeddings/private.bin", b"private")
        with self.assertRaisesRegex(ValueError, "嵌入附件"):
            clear_author_properties(embedded.getvalue(), ".docx")

    def test_clear_authors_preserves_contents_and_times(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = root / "sample.docx"
            document = Document()
            document.add_paragraph("正文署名：张三")
            document.core_properties.author = "原作者"
            document.core_properties.last_modified_by = "修改者"
            document.save(doc)
            xlsx = root / "sample.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "原作者"
            workbook.active["B1"] = "=1+1"
            workbook.active["A1"].comment = Comment("私人批注", "批注作者")
            workbook.properties.creator = "原作者"
            workbook.properties.lastModifiedBy = "修改者"
            workbook.save(xlsx)
            txt = root / "sample.txt"
            txt.write_text("正文", encoding="utf-8")
            bad = root / "broken.docx"
            bad.write_bytes(b"not a zip")
            for source in [doc, xlsx]:
                os.utime(source, ns=(1600000000123456700, 1600000000765432100))
            with patch.object(socket.socket, "connect", side_effect=AssertionError("network prohibited")):
                results = replace_batch([doc, txt, bad, xlsx], root / "out", None, mode="clear_authors")
            self.assertEqual([r["ok"] for r in results], [True, False, False, True])
            for result in (results[0], results[3]):
                self.assertGreater(result["count"], 2)
                source, output = Path(result["source"]), Path(result["output"])
                self.assertEqual(source.stat().st_ctime_ns, output.stat().st_ctime_ns)
                self.assertEqual(source.stat().st_mtime_ns, output.stat().st_mtime_ns)
                with ZipFile(source) as original, ZipFile(output) as cleaned:
                    self.assertFalse(any("comments" in name.lower() for name in cleaned.namelist()))
                    for name in cleaned.namelist():
                        self.assertEqual(original.getinfo(name).date_time, cleaned.getinfo(name).date_time)
                        if name in ("xl/styles.xml", "word/styles.xml"):
                            self.assertEqual(original.read(name), cleaned.read(name))
                    for name in ("docProps/core.xml", "docProps/app.xml"):
                        self.assertFalse(any(node.text or node.attrib for node in etree.fromstring(cleaned.read(name))))
                _, repeat = clear_author_properties(output.read_bytes(), output.suffix)
                self.assertEqual(repeat["count"], 0)
            clean_doc = Document(results[0]["output"])
            self.assertFalse(clean_doc.core_properties.author)
            self.assertFalse(clean_doc.core_properties.last_modified_by)
            self.assertEqual(clean_doc.paragraphs[0].text, "正文署名：张三")
            clean_book = load_workbook(results[3]["output"])
            self.assertFalse(clean_book.properties.creator)
            self.assertFalse(clean_book.properties.lastModifiedBy)
            self.assertEqual(clean_book.active["B1"].value, "=1+1")
            self.assertIsNone(clean_book.active["A1"].comment)
            self.assertEqual(Document(doc).core_properties.author, "原作者")
            repeated = replace_batch([doc], root / "out", None, mode="clear_authors")
            self.assertTrue(repeated[0]["output"].endswith("_cleaned_2.docx"))
            engine = SensitiveWordFilter([{"word": "张三", "replacement": "某人"}])
            for mode in ("replace", "restore"):
                reports = replace_batch([doc, xlsx], root / mode, engine, mode=mode)
                for report in reports:
                    self.assertTrue(report["ok"], report)
                    self.assertGreater(report["privacy"]["count"], 0)
                    with ZipFile(report["output"]) as archive:
                        self.assertFalse(any(node.text or node.attrib for node in etree.fromstring(archive.read("docProps/core.xml"))))
                        self.assertFalse(any("comments" in n.lower() for n in archive.namelist()))
                    before, after = Path(report["source"]).stat(), Path(report["output"]).stat()
                    self.assertEqual(before.st_ctime_ns, after.st_ctime_ns)
                    self.assertEqual(before.st_mtime_ns, after.st_mtime_ns)


if __name__ == "__main__":
    unittest.main()
