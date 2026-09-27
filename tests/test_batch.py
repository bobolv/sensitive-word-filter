import os
import json
import socket
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook

from sensitive_filter import SensitiveWordFilter
from sensitive_filter.batch import replace_batch, replace_file, load_batch_engine


@unittest.skipUnless(os.name == "nt", "Windows creation time semantics")
class BatchTests(unittest.TestCase):
    def test_mixed_offline_batch_and_exact_times(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            text = root / "输入.txt"
            text.write_bytes(b"\xef\xbb\xbfsecret\r\n")
            doc = root / "输入.docx"
            document = Document()
            document.add_paragraph("secret")
            document.save(doc)
            xlsx = root / "输入.xlsx"
            book = Workbook()
            book.active["A1"] = "secret"
            book.active["B1"] = "=1+1"
            book.save(xlsx)
            bad = root / "bad.txt"
            bad.write_bytes(b"\xff\xff")
            sources = [text, bad, doc, xlsx, text]
            for source in sources:
                os.utime(source, ns=(1600000000123456700, 1600000000765432100))
            before = {source: source.stat() for source in sources}
            engine = SensitiveWordFilter([{"word": "secret", "replacement": "safe"}])
            with patch.object(socket.socket, "connect", side_effect=AssertionError("network prohibited")):
                results = replace_batch(sources, root / "out", engine)
            self.assertEqual([r["ok"] for r in results], [True, False, True, True, True])
            self.assertEqual(len(list((root / "out").iterdir())), 4)
            for item in results:
                if item["ok"]:
                    self.assertEqual(item["changes"], [{"before": "secret", "after": "safe", "count": 1}])
            for result in results:
                if not result["ok"]:
                    continue
                source = Path(result["source"])
                target = Path(result["output"])
                self.assertEqual(target.stat().st_mtime_ns, before[source].st_mtime_ns)
                self.assertEqual(target.stat().st_ctime_ns, before[source].st_ctime_ns)
            self.assertEqual(Path(results[0]["output"]).read_bytes(), b"\xef\xbb\xbfsafe\r\n")
            self.assertEqual(Document(results[2]["output"]).paragraphs[0].text, "safe")
            filtered = load_workbook(results[3]["output"])
            self.assertEqual(filtered.active["A1"].value, "safe")
            self.assertEqual(filtered.active["B1"].value, "=1+1")
            self.assertEqual(text.read_bytes(), b"\xef\xbb\xbfsecret\r\n")
            with self.assertRaises(FileExistsError):
                replace_file(text, text, engine)
            failed_target = root / "failed.txt"
            with patch("sensitive_filter.batch.copy_times", side_effect=OSError("denied")):
                with self.assertRaises(OSError):
                    replace_file(text, failed_target, engine)
            self.assertFalse(failed_target.exists())

    def test_restore_mixed_files_and_actual_edit_details(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            words = root / "words.json"
            words.write_text(json.dumps({"words": [
                {"word": "张三", "replacement": "某人员"},
                {"word": "张", "replacement": "某"},
                {"word": "机密"},
            ]}), encoding="utf-8")
            forward = load_batch_engine(words)
            reverse = load_batch_engine(words, "restore")
            text = root / "a.txt"
            text.write_text("张三 张三 机密", encoding="utf-8")
            doc = root / "a.docx"
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run("张")
            paragraph.add_run("三")
            document.save(doc)
            xlsx = root / "a.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "张三"
            workbook.save(xlsx)
            with patch.object(socket.socket, "connect", side_effect=AssertionError("network prohibited")):
                replaced = replace_batch([text, doc, xlsx], root / "out", forward)
                restored = replace_batch([r["output"] for r in replaced], root / "back", reverse, mode="restore")
            self.assertTrue(all(r["ok"] for r in restored))
            self.assertEqual(replaced[0]["count"], 3)  # overlapping 张 is not an extra edit
            self.assertEqual(replaced[0]["changes"], [
                {"before": "张三", "after": "某人员", "count": 2},
                {"before": "机密", "after": "**", "count": 1},
            ])
            self.assertEqual(restored[0]["changes"], [{"before": "某人员", "after": "张三", "count": 2}])
            self.assertEqual(Path(restored[0]["output"]).read_text(encoding="utf-8"), "张三 张三 **")
            self.assertEqual(Document(restored[1]["output"]).paragraphs[0].text, "张三")
            self.assertEqual(load_workbook(restored[2]["output"]).active["A1"].value, "张三")
            for item in restored:
                self.assertIn("_restored", item["output"])
                original, output = Path(item["source"]).stat(), Path(item["output"]).stat()
                self.assertEqual(original.st_ctime_ns, output.st_ctime_ns)
                self.assertEqual(original.st_mtime_ns, output.st_mtime_ns)
            words.write_text(json.dumps([{"word": "a", "replacement": "same"},
                                         {"word": "b", "replacement": "SAME"}]), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "无法唯一恢复"):
                load_batch_engine(words, "restore")
            words.write_text('[{"word": "a"}]', encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "没有可反向恢复"):
                load_batch_engine(words, "restore")


if __name__ == "__main__":
    unittest.main()
