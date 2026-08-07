from io import BytesIO
import base64

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from sensitive_filter import SensitiveWordFilter
from sensitive_filter.docx_handler import (
    create_writing_template,
    document_text,
    docx_to_markdown,
    replace_docx,
    scan_docx,
)


engine = SensitiveWordFilter([{"word": "示例敏感词", "category": "测试", "level": 1}])
document = Document()
paragraph = document.add_paragraph("正文：")
paragraph.add_run("示例").bold = True
paragraph.add_run("敏感词")

# 多行、合并和嵌套表格覆盖曾经因对象 ID 复用而漏读的场景。
table = document.add_table(rows=120, cols=2)
expected_table_matches = 0
for index, row in enumerate(table.rows):
    row.cells[0].text = f"第 {index} 行"
    row.cells[1].text = f"表格示例敏感词 {index}"
    expected_table_matches += 1
merged = table.cell(0, 0).merge(table.cell(0, 1))
merged.text = "合并单元格示例敏感词"
nested = table.cell(1, 0).add_table(rows=1, cols=1)
nested.cell(0, 0).text = "嵌套表示例敏感词"
expected_table_matches += 1

document.sections[0].header.paragraphs[0].text = "页眉示例敏感词"
document.sections[0].footer.paragraphs[0].text = "页脚示例敏感词"
source = BytesIO()
document.save(source)

expected_total = 1 + expected_table_matches + 2
scanned = scan_docx(source.getvalue(), engine)
assert scanned["count"] == expected_total, (scanned["count"], expected_total)
assert scanned["text_length"] > 1_000

output, result = replace_docx(source.getvalue(), engine)
assert result["count"] == expected_total
filtered = Document(BytesIO(output))
filtered_text = document_text(filtered)
assert "示例敏感词" not in filtered_text
assert filtered_text.count("*****") == expected_total
assert filtered.paragraphs[0].runs[1].bold is True

# 页面只返回有限预览，但文末命中仍必须被检测和替换。
large_document = Document()
large_document.add_paragraph("普通内容" * 55_000)
large_document.add_paragraph("文末示例敏感词")
large_source = BytesIO()
large_document.save(large_source)
large_scan = scan_docx(large_source.getvalue(), engine)
assert large_scan["text_length"] > 200_000
assert large_scan["text_truncated"] is True
assert large_scan["count"] == 1
large_output, large_result = replace_docx(large_source.getvalue(), engine)
assert large_result["count"] == 1
assert "示例敏感词" not in document_text(Document(BytesIO(large_output)))

# 替换后转换 Markdown：保留标题层级和表格，忽略目录及图片。
structured = Document()
visual_title = structured.add_paragraph("未使用标题样式的项目方案")
visual_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
visual_title_run = visual_title.runs[0]
visual_title_run.bold = True
visual_title_run.font.size = Pt(20)
structured.add_paragraph("第一章 示例敏感词", style="Heading 1")
structured.add_paragraph("第一节 说明", style="Heading 2")
structured.add_paragraph("正文示例敏感词")
custom_heading = structured.styles.add_style("业务三级标题", WD_STYLE_TYPE.PARAGRAPH)
custom_heading.base_style = structured.styles["Heading 3"]
structured.add_paragraph("1.1.1 自定义样式标题", style=custom_heading)
structured.add_paragraph("自定义标题下的正文")
structured.add_paragraph("一、手工编号标题").runs[0].bold = True
structured.add_paragraph("手工标题下的正文")
toc_style = structured.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
structured.add_paragraph("目录中的项目", style=toc_style)
structured_table = structured.add_table(rows=2, cols=2)
structured_table.cell(0, 0).text = "名称"
structured_table.cell(0, 1).text = "说明"
structured_table.cell(1, 0).text = "项目|一"
structured_table.cell(1, 1).text = "表格示例敏感词"
structured.add_paragraph("两张表之间会被删除的正文")
second_table = structured.add_table(rows=2, cols=1)
second_table.cell(0, 0).text = "第二张表表头"
second_table.cell(1, 0).text = "第二张表数据"
pixel_png = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)
structured.add_paragraph().add_run().add_picture(BytesIO(pixel_png), width=Inches(0.1))
structured.add_paragraph("图1 系统结构", style="Caption")
structured_source = BytesIO()
structured.save(structured_source)

markdown_output, markdown_result = docx_to_markdown(structured_source.getvalue(), engine)
markdown = markdown_output.decode("utf-8-sig")
assert "# 第一章 *****" in markdown
assert "## 第一节 说明" in markdown
assert "### 1.1.1 自定义样式标题" in markdown
assert "正文*****" in markdown
assert "| 名称 | 说明 |" in markdown
assert "| 项目\\|一 | 表格***** |" in markdown
assert "目录中的项目" not in markdown
assert markdown_result["heading_count"] == 3
assert markdown_result["table_count"] == 2

# 编写模板：保留标题、表头、表格结构和图片题注，清空正文、图片及数据行内容。
template_output, template_result = create_writing_template(structured_source.getvalue())
template = Document(BytesIO(template_output))
template_text = document_text(template)
assert "第一章 示例敏感词" in template_text
assert "第一节 说明" in template_text
assert "1.1.1 自定义样式标题" in template_text
assert "未使用标题样式的项目方案" in template_text
assert "一、手工编号标题" in template_text
assert "[请" not in template_text
assert "正文示例敏感词" not in template_text
assert "自定义标题下的正文" not in template_text
assert "手工标题下的正文" not in template_text
assert "两张表之间会被删除的正文" not in template_text
assert "目录中的项目" not in template_text
assert "图1 系统结构" in template_text
assert [cell.text for cell in template.tables[0].rows[0].cells] == ["名称", "说明"]
assert all(not cell.text for row in template.tables[0].rows[1:] for cell in row.cells)
assert template.tables[1].cell(0, 0).text == "第二张表表头"
assert template.tables[1].cell(1, 0).text == ""
for table in template.tables:
    next_element = table._tbl.getnext()
    assert next_element is not None and next_element.tag == qn("w:p")
assert not template.inline_shapes
assert template_result == {"heading_count": 5, "table_count": 2, "caption_count": 1}
print("DOCX paragraphs/tables/merged/nested scan and replace passed")
